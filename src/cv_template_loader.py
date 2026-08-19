"""Load and validate registered CV templates and reference documents.

Phase 5.1 rules:
- career_data/ remains the factual source of truth.
- CV files may provide structure and approved wording references only.
- Only files registered in source_documents/cv/cv_index.yaml are loadable.
- Excluded files must never enter the tailoring pipeline.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from docx import Document


DEFAULT_INDEX_PATH = Path("source_documents/cv/cv_index.yaml")


class CVTemplateLoaderError(RuntimeError):
    """Raised when a CV registry or registered document cannot be loaded safely."""


@dataclass(slots=True)
class CVRegistryEntry:
    """One document entry from cv_index.yaml."""

    key: str
    file: str
    document_type: str
    status: str
    role_families: list[str] = field(default_factory=list)
    uses: list[str] = field(default_factory=list)
    wording_source: bool = False
    factual_authority: bool = False
    priority: int | None = None
    notes: str | None = None


@dataclass(slots=True)
class LoadedCVDocument:
    """Text and metadata extracted from one registered DOCX document."""

    registry_key: str
    path: Path
    document_type: str
    status: str
    role_families: list[str]
    uses: list[str]
    wording_source: bool
    factual_authority: bool
    priority: int | None
    paragraphs: list[str]
    table_rows: list[list[str]]
    full_text: str
    notes: str | None = None


@dataclass(slots=True)
class CVRegistry:
    """Validated representation of source_documents/cv/cv_index.yaml."""

    version: int
    factual_source_of_truth: str
    policy: dict[str, Any]
    documents: dict[str, CVRegistryEntry]
    selection_rules: dict[str, Any]
    excluded_files: set[str]
    index_path: Path

    @property
    def cv_directory(self) -> Path:
        return self.index_path.parent

    @property
    def active_documents(self) -> list[CVRegistryEntry]:
        return [
            entry
            for entry in self.documents.values()
            if entry.status.casefold() == "active"
        ]

    @property
    def master_templates(self) -> list[CVRegistryEntry]:
        return [
            entry
            for entry in self.active_documents
            if entry.document_type == "master_template"
        ]

    @property
    def reference_cvs(self) -> list[CVRegistryEntry]:
        return [
            entry
            for entry in self.active_documents
            if entry.document_type == "reference_cv"
        ]


def _repository_root(start: Path | None = None) -> Path:
    """Find the repository root by walking upward from start or CWD."""

    current = (start or Path.cwd()).expanduser().resolve()

    if current.is_file():
        current = current.parent

    for candidate in [current, *current.parents]:
        if (candidate / "career_data").exists() and (candidate / "src").exists():
            return candidate

    raise CVTemplateLoaderError(
        "Could not locate repository root containing career_data/ and src/."
    )


def _resolve_index_path(
    index_path: str | Path | None = None,
    repository_root: str | Path | None = None,
) -> Path:
    if index_path is not None:
        resolved = Path(index_path).expanduser()
        if not resolved.is_absolute():
            root = (
                Path(repository_root).expanduser().resolve()
                if repository_root
                else _repository_root()
            )
            resolved = root / resolved
        return resolved.resolve()

    root = (
        Path(repository_root).expanduser().resolve()
        if repository_root
        else _repository_root()
    )
    return (root / DEFAULT_INDEX_PATH).resolve()


def _require_mapping(value: Any, label: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise CVTemplateLoaderError(f"{label} must be a YAML mapping.")
    return value


def _require_string(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise CVTemplateLoaderError(f"{label} must be a non-empty string.")
    return value.strip()


def load_cv_registry(
    index_path: str | Path | None = None,
    *,
    repository_root: str | Path | None = None,
) -> CVRegistry:
    """Load and validate the CV registry."""

    path = _resolve_index_path(index_path, repository_root)

    if not path.exists():
        raise CVTemplateLoaderError(f"CV index not found: {path}")

    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    except yaml.YAMLError as exc:
        raise CVTemplateLoaderError(
            f"Invalid YAML in CV index: {path}"
        ) from exc

    data = _require_mapping(raw, "cv_index.yaml")

    version = data.get("version")
    if not isinstance(version, int) or version < 1:
        raise CVTemplateLoaderError("CV index version must be an integer >= 1.")

    factual_source = _require_string(
        data.get("factual_source_of_truth"),
        "factual_source_of_truth",
    )

    policy = _require_mapping(data.get("policy", {}), "policy")
    selection_rules = _require_mapping(
        data.get("selection_rules", {}),
        "selection_rules",
    )

    excluded_raw = data.get("excluded_files", [])
    if not isinstance(excluded_raw, list):
        raise CVTemplateLoaderError("excluded_files must be a list.")

    excluded_files = {
        _require_string(item, "excluded_files item")
        for item in excluded_raw
    }

    documents_raw = _require_mapping(data.get("documents"), "documents")
    documents: dict[str, CVRegistryEntry] = {}

    for key, raw_entry in documents_raw.items():
        entry_data = _require_mapping(
            raw_entry,
            f"documents.{key}",
        )

        file_name = _require_string(
            entry_data.get("file"),
            f"documents.{key}.file",
        )

        if file_name in excluded_files:
            raise CVTemplateLoaderError(
                f"Registered document {key!r} is also excluded: {file_name}"
            )

        role_families = entry_data.get("role_families", [])
        uses = entry_data.get("uses", [])

        if not isinstance(role_families, list):
            raise CVTemplateLoaderError(
                f"documents.{key}.role_families must be a list."
            )

        if not isinstance(uses, list):
            raise CVTemplateLoaderError(
                f"documents.{key}.uses must be a list."
            )

        documents[str(key)] = CVRegistryEntry(
            key=str(key),
            file=file_name,
            document_type=_require_string(
                entry_data.get("document_type"),
                f"documents.{key}.document_type",
            ),
            status=_require_string(
                entry_data.get("status"),
                f"documents.{key}.status",
            ),
            role_families=[
                _require_string(
                    item,
                    f"documents.{key}.role_families item",
                )
                for item in role_families
            ],
            uses=[
                _require_string(
                    item,
                    f"documents.{key}.uses item",
                )
                for item in uses
            ],
            wording_source=bool(entry_data.get("wording_source", False)),
            factual_authority=bool(
                entry_data.get("factual_authority", False)
            ),
            priority=(
                int(entry_data["priority"])
                if entry_data.get("priority") is not None
                else None
            ),
            notes=(
                str(entry_data["notes"]).strip()
                if entry_data.get("notes")
                else None
            ),
        )

    registry = CVRegistry(
        version=version,
        factual_source_of_truth=factual_source,
        policy=policy,
        documents=documents,
        selection_rules=selection_rules,
        excluded_files=excluded_files,
        index_path=path,
    )

    validate_cv_registry_files(registry)
    return registry


def validate_cv_registry_files(registry: CVRegistry) -> None:
    """Ensure all active registered files exist and stay inside the CV folder."""

    cv_dir = registry.cv_directory.resolve()

    for entry in registry.active_documents:
        path = (cv_dir / entry.file).resolve()

        try:
            path.relative_to(cv_dir)
        except ValueError as exc:
            raise CVTemplateLoaderError(
                f"Registered CV path escapes CV directory: {entry.file}"
            ) from exc

        if not path.exists():
            raise CVTemplateLoaderError(
                f"Registered CV file does not exist: {path}"
            )

        if path.suffix.casefold() != ".docx":
            raise CVTemplateLoaderError(
                f"Registered CV must be a .docx file: {path}"
            )

    default_master = registry.selection_rules.get(
        "default_master_template"
    )
    if default_master:
        if default_master not in registry.documents:
            raise CVTemplateLoaderError(
                "selection_rules.default_master_template refers to an "
                f"unknown registry key: {default_master}"
            )

        entry = registry.documents[default_master]
        if entry.document_type != "master_template":
            raise CVTemplateLoaderError(
                "Default master template must have document_type "
                "'master_template'."
            )


def get_registry_entry(
    registry: CVRegistry,
    registry_key: str,
) -> CVRegistryEntry:
    """Return one registered CV entry by key."""

    try:
        return registry.documents[registry_key]
    except KeyError as exc:
        raise CVTemplateLoaderError(
            f"Unknown CV registry key: {registry_key}"
        ) from exc


def resolve_cv_path(
    registry: CVRegistry,
    registry_key: str,
) -> Path:
    """Resolve a registered CV path without allowing arbitrary file access."""

    entry = get_registry_entry(registry, registry_key)

    if entry.file in registry.excluded_files:
        raise CVTemplateLoaderError(
            f"CV file is explicitly excluded: {entry.file}"
        )

    path = (registry.cv_directory / entry.file).resolve()
    cv_dir = registry.cv_directory.resolve()

    try:
        path.relative_to(cv_dir)
    except ValueError as exc:
        raise CVTemplateLoaderError(
            f"CV path escapes registered CV directory: {entry.file}"
        ) from exc

    if not path.exists():
        raise CVTemplateLoaderError(f"CV file not found: {path}")

    return path


def _extract_docx_content(
    path: Path,
) -> tuple[list[str], list[list[str]], str]:
    """Extract paragraph and table text while preserving document order loosely."""

    try:
        document = Document(path)
    except Exception as exc:
        raise CVTemplateLoaderError(
            f"Could not open DOCX file: {path}"
        ) from exc

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    table_rows: list[list[str]] = []
    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            if any(values):
                table_rows.append(values)

    table_text = [
        " | ".join(cell for cell in row if cell)
        for row in table_rows
        if any(cell for cell in row)
    ]

    full_text = "\n".join([*paragraphs, *table_text]).strip()

    return paragraphs, table_rows, full_text


def load_cv_document(
    registry: CVRegistry,
    registry_key: str,
) -> LoadedCVDocument:
    """Load one registered DOCX document and its registry metadata."""

    entry = get_registry_entry(registry, registry_key)

    if entry.status.casefold() != "active":
        raise CVTemplateLoaderError(
            f"CV document is not active: {registry_key}"
        )

    path = resolve_cv_path(registry, registry_key)
    paragraphs, table_rows, full_text = _extract_docx_content(path)

    if not full_text:
        raise CVTemplateLoaderError(
            f"Registered CV contains no readable text: {path}"
        )

    return LoadedCVDocument(
        registry_key=entry.key,
        path=path,
        document_type=entry.document_type,
        status=entry.status,
        role_families=list(entry.role_families),
        uses=list(entry.uses),
        wording_source=entry.wording_source,
        factual_authority=entry.factual_authority,
        priority=entry.priority,
        paragraphs=paragraphs,
        table_rows=table_rows,
        full_text=full_text,
        notes=entry.notes,
    )


def load_default_master_template(
    registry: CVRegistry,
) -> LoadedCVDocument:
    """Load the default structural master configured in cv_index.yaml."""

    key = registry.selection_rules.get("default_master_template")
    if not isinstance(key, str) or not key.strip():
        raise CVTemplateLoaderError(
            "No default_master_template is configured."
        )

    document = load_cv_document(registry, key)

    if document.document_type != "master_template":
        raise CVTemplateLoaderError(
            "Configured default master is not a master_template."
        )

    return document


def load_active_reference_cvs(
    registry: CVRegistry,
) -> list[LoadedCVDocument]:
    """Load all active reference CVs, ordered by priority then registry key."""

    entries = sorted(
        registry.reference_cvs,
        key=lambda item: (
            item.priority if item.priority is not None else 999,
            item.key,
        ),
    )

    return [
        load_cv_document(registry, entry.key)
        for entry in entries
    ]
