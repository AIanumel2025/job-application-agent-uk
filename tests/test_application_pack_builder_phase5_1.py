"""Phase 5.1 integration tests for the application-pack builder."""

from pathlib import Path

import yaml
from docx import Document

import src.cv_tailor as cv_tailor_module
from src.application_pack_builder import build_application_pack
from src.cv_template_loader import load_cv_registry
from tests.phase5_helpers import make_match, make_profile


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _registry(tmp_path: Path):
    cv_dir = tmp_path / "source_documents" / "cv"

    _write_docx(
        cv_dir / "master.docx",
        [
            "ANTHONY L. ANUMEL",
            ">>> SWAP THIS BLOCK FOR EVERY APPLICATION — TAILOR IN UNDER 2 MINUTES <<<",
            "[TITLE EXACTLY AS WRITTEN IN THE JOB POSTING]",
            "PROFESSIONAL SNAPSHOT",
            "CORE COMPETENCIES",
            ">>> KEEP STABLE BELOW THIS LINE — DO NOT EDIT FOR EACH APPLICATION <<<",
            "TECHNICAL SKILLS",
            "PROFESSIONAL EXPERIENCE",
        ],
    )

    _write_docx(
        cv_dir / "data_reference.docx",
        [
            "PROFESSIONAL EXPERIENCE",
            (
                "Used Excel and Power BI dashboards to improve student "
                "performance and communicate insights."
            ),
            "KEY PROJECTS",
            (
                "Built a Python and AWS document-processing pipeline "
                "using Docker and GitHub Actions."
            ),
        ],
    )

    payload = {
        "version": 1,
        "factual_source_of_truth": "career_data/",
        "policy": {
            "cv_documents_are_factual_authority": False,
        },
        "documents": {
            "master": {
                "file": "master.docx",
                "document_type": "master_template",
                "status": "active",
                "role_families": [
                    "data_engineering",
                    "ai_engineering",
                ],
                "uses": ["document_structure"],
                "wording_source": False,
                "factual_authority": False,
            },
            "data_reference": {
                "file": "data_reference.docx",
                "document_type": "reference_cv",
                "status": "active",
                "priority": 1,
                "role_families": [
                    "data_engineering",
                    "data_analytics",
                ],
                "uses": ["approved_wording_reference"],
                "wording_source": True,
                "factual_authority": False,
            },
        },
        "selection_rules": {
            "default_master_template": "master",
            "maximum_reference_cvs_per_application": 2,
        },
        "excluded_files": [],
    }

    index = cv_dir / "cv_index.yaml"
    index.write_text(
        yaml.safe_dump(payload, sort_keys=False),
        encoding="utf-8",
    )

    return load_cv_registry(index)


def test_application_pack_uses_phase_5_1_template_flow(
    tmp_path,
    monkeypatch,
) -> None:
    registry = _registry(tmp_path)

    original_load = cv_tailor_module.load_cv_registry
    monkeypatch.setattr(
        cv_tailor_module,
        "load_cv_registry",
        lambda: registry,
    )

    try:
        pack = build_application_pack(
            make_profile(),
            make_match(),
        )
    finally:
        monkeypatch.setattr(
            cv_tailor_module,
            "load_cv_registry",
            original_load,
        )

    assert pack.cv.target_role == "Data Engineer"
    assert "# Anthony L. Anumel" in pack.cv.markdown
    assert "## Data Engineer" in pack.cv.markdown
    assert (
        pack.metadata.get("cv_template_integration")
        == "phase_5_1"
    )


def test_application_pack_still_requires_human_review(
    tmp_path,
    monkeypatch,
) -> None:
    registry = _registry(tmp_path)

    original_load = cv_tailor_module.load_cv_registry
    monkeypatch.setattr(
        cv_tailor_module,
        "load_cv_registry",
        lambda: registry,
    )

    try:
        pack = build_application_pack(
            make_profile(),
            make_match(),
        )
    finally:
        monkeypatch.setattr(
            cv_tailor_module,
            "load_cv_registry",
            original_load,
        )

    assert pack.manifest.human_review_required is True


def test_application_pack_does_not_leak_template_placeholders(
    tmp_path,
    monkeypatch,
) -> None:
    registry = _registry(tmp_path)

    original_load = cv_tailor_module.load_cv_registry
    monkeypatch.setattr(
        cv_tailor_module,
        "load_cv_registry",
        lambda: registry,
    )

    try:
        pack = build_application_pack(
            make_profile(),
            make_match(),
        )
    finally:
        monkeypatch.setattr(
            cv_tailor_module,
            "load_cv_registry",
            original_load,
        )

    markdown = pack.cv.markdown

    assert "[TITLE EXACTLY AS WRITTEN" not in markdown
    assert "SWAP THIS BLOCK" not in markdown
    assert "KEEP STABLE BELOW THIS LINE" not in markdown
