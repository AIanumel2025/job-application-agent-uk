"""Phase 5.1 integration tests for template-aware CV tailoring."""

from pathlib import Path

import yaml
from docx import Document

from src.cv_template_loader import load_cv_registry
from src.cv_tailor import tailor_cv
from tests.phase5_helpers import (
    make_match,
    make_profile,
    make_selected_evidence,
)


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _build_registry(tmp_path: Path):
    cv_dir = tmp_path / "source_documents" / "cv"

    _write_docx(
        cv_dir / "master.docx",
        [
            "ANTHONY L. ANUMEL",
            ">>> SWAP THIS BLOCK FOR EVERY APPLICATION — TAILOR IN UNDER 2 MINUTES <<<",
            "[TITLE EXACTLY AS WRITTEN IN THE JOB POSTING]",
            "PROFESSIONAL SNAPSHOT",
            "[Line 1: Mirror the job description]",
            "CORE COMPETENCIES",
            ">>> KEEP STABLE BELOW THIS LINE — DO NOT EDIT FOR EACH APPLICATION <<<",
            "TECHNICAL SKILLS",
            "PROFESSIONAL EXPERIENCE",
            "[Your Title] | [Company]",
        ],
    )

    _write_docx(
        cv_dir / "data_reference.docx",
        [
            "PROFESSIONAL EXPERIENCE",
            (
                "Used Excel and Power BI dashboards to improve student "
                "performance and communicate insights."
            ),
            "KEY PROJECTS",
            (
                "Built a Python and AWS document-processing pipeline "
                "using Docker and GitHub Actions."
            ),
        ],
    )

    payload = {
        "version": 1,
        "factual_source_of_truth": "career_data/",
        "policy": {
            "cv_documents_are_factual_authority": False,
            "allow_reference_wording": True,
            "require_claim_verification_against_career_data": True,
            "allow_unverified_metrics": False,
            "allow_confidential_content": False,
            "human_review_required_before_submission": True,
        },
        "documents": {
            "master": {
                "file": "master.docx",
                "document_type": "master_template",
                "status": "active",
                "role_families": [
                    "data_engineering",
                    "ai_engineering",
                ],
                "uses": [
                    "document_structure",
                    "section_order",
                ],
                "wording_source": False,
                "factual_authority": False,
            },
            "data_reference": {
                "file": "data_reference.docx",
                "document_type": "reference_cv",
                "status": "active",
                "priority": 1,
                "role_families": [
                    "data_engineering",
                    "data_analytics",
                ],
                "uses": [
                    "approved_wording_reference",
                    "experience_bullet_reference",
                    "project_wording_reference",
                ],
                "wording_source": True,
                "factual_authority": False,
            },
        },
        "selection_rules": {
            "default_master_template": "master",
            "prefer_role_family_match": True,
            "prefer_priority_1_reference": True,
            "allow_multiple_reference_cvs": True,
            "maximum_reference_cvs_per_application": 2,
        },
        "excluded_files": [],
    }

    index = cv_dir / "cv_index.yaml"
    index.write_text(
        yaml.safe_dump(payload, sort_keys=False),
        encoding="utf-8",
    )

    return load_cv_registry(index)


def test_template_aware_tailor_keeps_candidate_and_role(tmp_path) -> None:
    registry = _build_registry(tmp_path)

    cv = tailor_cv(
        make_profile(),
        make_match(),
        make_selected_evidence(),
        registry=registry,
    )

    assert cv.candidate_name == "Anthony L. Anumel"
    assert cv.target_role == "Data Engineer"
    assert "# Anthony L. Anumel" in cv.markdown
    assert "## Data Engineer" in cv.markdown


def test_template_aware_tailor_can_reuse_verified_reference_wording(
    tmp_path,
) -> None:
    registry = _build_registry(tmp_path)

    cv = tailor_cv(
        make_profile(),
        make_match(),
        make_selected_evidence(),
        registry=registry,
    )

    assert (
        "Built a Python and AWS document-processing pipeline "
        "using Docker and GitHub Actions."
        in cv.markdown
    )


def test_template_placeholders_do_not_leak_into_tailored_cv(
    tmp_path,
) -> None:
    registry = _build_registry(tmp_path)

    cv = tailor_cv(
        make_profile(),
        make_match(),
        make_selected_evidence(),
        registry=registry,
    )

    assert "[TITLE EXACTLY AS WRITTEN" not in cv.markdown
    assert "[Your Title]" not in cv.markdown
    assert "SWAP THIS BLOCK" not in cv.markdown
    assert "KEEP STABLE BELOW THIS LINE" not in cv.markdown


def test_template_selection_reason_is_preserved_as_warning(
    tmp_path,
) -> None:
    registry = _build_registry(tmp_path)

    cv = tailor_cv(
        make_profile(),
        make_match(),
        make_selected_evidence(),
        registry=registry,
    )

    assert cv.warnings
    assert any(
        "reference" in warning.casefold()
        or "role family" in warning.casefold()
        for warning in cv.warnings
    )
